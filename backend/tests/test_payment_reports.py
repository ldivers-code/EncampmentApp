"""
Payment Report Import and Summary Tests
Tests for eCAP payment report import, payment summary, and import history endpoints
"""
import pytest
import requests
import os
import io

BASE_URL = os.environ.get('REACT_APP_BACKEND_URL', '').rstrip('/')

# Test credentials
COMMANDER_EMAIL = "commander@test.com"
COMMANDER_PASSWORD = "test123"

# Non-finance user for access denial tests
CADRE_EMAIL = "cadre@cap.us"
CADRE_PASSWORD = "test123"


@pytest.fixture(scope="module")
def commander_token():
    """Get commander auth token (has finance access)"""
    response = requests.post(f"{BASE_URL}/api/auth/login", json={
        "email": COMMANDER_EMAIL,
        "password": COMMANDER_PASSWORD
    })
    if response.status_code != 200:
        pytest.skip(f"Commander login failed: {response.status_code}")
    return response.json().get("access_token")


@pytest.fixture(scope="module")
def commander_headers(commander_token):
    """Headers with commander auth"""
    return {"Authorization": f"Bearer {commander_token}"}


@pytest.fixture(scope="module")
def cadre_token():
    """Get cadre auth token (no finance access)"""
    response = requests.post(f"{BASE_URL}/api/auth/login", json={
        "email": CADRE_EMAIL,
        "password": CADRE_PASSWORD
    })
    if response.status_code != 200:
        pytest.skip(f"Cadre login failed: {response.status_code}")
    return response.json().get("access_token")


@pytest.fixture(scope="module")
def cadre_headers(cadre_token):
    """Headers with cadre auth"""
    return {"Authorization": f"Bearer {cadre_token}"}


class TestPaymentSummaryEndpoint:
    """Tests for GET /api/participants/payment-summary"""
    
    def test_payment_summary_returns_200(self, commander_headers):
        """Payment summary endpoint returns 200 for finance users"""
        response = requests.get(f"{BASE_URL}/api/participants/payment-summary", headers=commander_headers)
        assert response.status_code == 200, f"Expected 200, got {response.status_code}: {response.text}"
    
    def test_payment_summary_structure(self, commander_headers):
        """Payment summary has correct structure"""
        response = requests.get(f"{BASE_URL}/api/participants/payment-summary", headers=commander_headers)
        assert response.status_code == 200
        data = response.json()
        
        # Check required top-level fields
        assert "total" in data, "Missing 'total' field"
        assert "paid" in data, "Missing 'paid' field"
        assert "unpaid" in data, "Missing 'unpaid' field"
        assert "by_type" in data, "Missing 'by_type' field"
        assert "by_flight" in data, "Missing 'by_flight' field"
        assert "by_status" in data, "Missing 'by_status' field"
        assert "unit_approved" in data, "Missing 'unit_approved' field"
        assert "wing_approved" in data, "Missing 'wing_approved' field"
        assert "participants" in data, "Missing 'participants' field"
        assert "last_import" in data, "Missing 'last_import' field"
    
    def test_payment_summary_totals_consistent(self, commander_headers):
        """Paid + unpaid should equal total"""
        response = requests.get(f"{BASE_URL}/api/participants/payment-summary", headers=commander_headers)
        assert response.status_code == 200
        data = response.json()
        
        assert data["paid"] + data["unpaid"] == data["total"], \
            f"Paid ({data['paid']}) + Unpaid ({data['unpaid']}) != Total ({data['total']})"
    
    def test_payment_summary_by_type_breakdown(self, commander_headers):
        """By type breakdown has correct structure"""
        response = requests.get(f"{BASE_URL}/api/participants/payment-summary", headers=commander_headers)
        assert response.status_code == 200
        data = response.json()
        
        for ptype, stats in data["by_type"].items():
            assert "total" in stats, f"Missing 'total' in by_type[{ptype}]"
            assert "paid" in stats, f"Missing 'paid' in by_type[{ptype}]"
            assert "unpaid" in stats, f"Missing 'unpaid' in by_type[{ptype}]"
            assert stats["paid"] + stats["unpaid"] == stats["total"], \
                f"by_type[{ptype}]: paid + unpaid != total"
    
    def test_payment_summary_by_flight_breakdown(self, commander_headers):
        """By flight breakdown has correct structure"""
        response = requests.get(f"{BASE_URL}/api/participants/payment-summary", headers=commander_headers)
        assert response.status_code == 200
        data = response.json()
        
        for flight, stats in data["by_flight"].items():
            assert "total" in stats, f"Missing 'total' in by_flight[{flight}]"
            assert "paid" in stats, f"Missing 'paid' in by_flight[{flight}]"
            assert "unpaid" in stats, f"Missing 'unpaid' in by_flight[{flight}]"
    
    def test_payment_summary_participants_list(self, commander_headers):
        """Participants list has correct structure"""
        response = requests.get(f"{BASE_URL}/api/participants/payment-summary", headers=commander_headers)
        assert response.status_code == 200
        data = response.json()
        
        assert isinstance(data["participants"], list), "participants should be a list"
        
        if len(data["participants"]) > 0:
            p = data["participants"][0]
            assert "capid" in p, "Missing 'capid' in participant"
            assert "name" in p, "Missing 'name' in participant"
            assert "participant_type" in p, "Missing 'participant_type' in participant"
            assert "flight" in p, "Missing 'flight' in participant"
            assert "paid" in p, "Missing 'paid' in participant"
            assert "unit_approved" in p, "Missing 'unit_approved' in participant"
            assert "wing_approved" in p, "Missing 'wing_approved' in participant"
    
    def test_payment_summary_denied_for_non_finance(self, cadre_headers):
        """Payment summary denied for non-finance users"""
        response = requests.get(f"{BASE_URL}/api/participants/payment-summary", headers=cadre_headers)
        assert response.status_code == 403, f"Expected 403, got {response.status_code}"


class TestPaymentImportHistoryEndpoint:
    """Tests for GET /api/payment-imports"""
    
    def test_payment_imports_returns_200(self, commander_headers):
        """Payment imports endpoint returns 200"""
        response = requests.get(f"{BASE_URL}/api/payment-imports", headers=commander_headers)
        assert response.status_code == 200, f"Expected 200, got {response.status_code}: {response.text}"
    
    def test_payment_imports_returns_list(self, commander_headers):
        """Payment imports returns a list"""
        response = requests.get(f"{BASE_URL}/api/payment-imports", headers=commander_headers)
        assert response.status_code == 200
        data = response.json()
        assert isinstance(data, list), "Response should be a list"
    
    def test_payment_imports_structure(self, commander_headers):
        """Payment import records have correct structure"""
        response = requests.get(f"{BASE_URL}/api/payment-imports", headers=commander_headers)
        assert response.status_code == 200
        data = response.json()
        
        if len(data) > 0:
            imp = data[0]
            assert "id" in imp, "Missing 'id' in import record"
            assert "filename" in imp, "Missing 'filename' in import record"
            assert "imported_by" in imp, "Missing 'imported_by' in import record"
            assert "imported_at" in imp, "Missing 'imported_at' in import record"
            assert "matched" in imp, "Missing 'matched' in import record"
            assert "updated" in imp, "Missing 'updated' in import record"
    
    def test_payment_imports_denied_for_non_finance(self, cadre_headers):
        """Payment imports denied for non-finance users"""
        response = requests.get(f"{BASE_URL}/api/payment-imports", headers=cadre_headers)
        assert response.status_code == 403, f"Expected 403, got {response.status_code}"


class TestPaymentImportEndpoint:
    """Tests for POST /api/participants/import-payments"""
    
    def test_import_rejects_non_excel(self, commander_headers):
        """Import rejects non-Excel files"""
        files = {"file": ("test.txt", io.BytesIO(b"test content"), "text/plain")}
        response = requests.post(
            f"{BASE_URL}/api/participants/import-payments",
            headers=commander_headers,
            files=files
        )
        assert response.status_code == 400, f"Expected 400, got {response.status_code}"
        assert "Excel" in response.json().get("detail", ""), "Error should mention Excel"
    
    def test_import_denied_for_non_finance(self, cadre_headers):
        """Import denied for non-finance users"""
        # Create a minimal xlsx file
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["RegistrantsCAPID", "PaidInFull", "AmountPaid"])
        ws.append(["123456", "Yes", "250.00"])
        
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        files = {"file": ("test.xlsx", buffer, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        response = requests.post(
            f"{BASE_URL}/api/participants/import-payments",
            headers=cadre_headers,
            files=files
        )
        assert response.status_code == 403, f"Expected 403, got {response.status_code}"
    
    def test_import_valid_excel_returns_200(self, commander_headers):
        """Import valid Excel file returns 200"""
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["RegistrantsCAPID", "PaidInFull", "AmountPaid", "RegistrationStatus", "UnitApproved", "WingApproved"])
        # Use a CAPID that likely doesn't exist to test not_found handling
        ws.append(["999999", "Yes", "250.00", "Approved", "Yes", "Yes"])
        
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        files = {"file": ("test_payment.xlsx", buffer, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        response = requests.post(
            f"{BASE_URL}/api/participants/import-payments",
            headers=commander_headers,
            files=files
        )
        assert response.status_code == 200, f"Expected 200, got {response.status_code}: {response.text}"
        
        data = response.json()
        assert "matched" in data, "Response should have 'matched' count"
        assert "updated" in data, "Response should have 'updated' count"
        assert "not_found" in data, "Response should have 'not_found' list"
        assert "import_id" in data, "Response should have 'import_id'"
    
    def test_import_creates_history_record(self, commander_headers):
        """Import creates a history record"""
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["RegistrantsCAPID", "PaidInFull"])
        ws.append(["888888", "No"])
        
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        # Get current import count
        history_before = requests.get(f"{BASE_URL}/api/payment-imports", headers=commander_headers)
        count_before = len(history_before.json())
        
        # Do import
        files = {"file": ("history_test.xlsx", buffer, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        response = requests.post(
            f"{BASE_URL}/api/participants/import-payments",
            headers=commander_headers,
            files=files
        )
        assert response.status_code == 200
        
        # Check history increased
        history_after = requests.get(f"{BASE_URL}/api/payment-imports", headers=commander_headers)
        count_after = len(history_after.json())
        
        assert count_after > count_before, "Import should create a history record"


class TestPaymentDataPersistence:
    """Tests for payment data persistence after import"""
    
    def test_payment_summary_reflects_import_data(self, commander_headers):
        """Payment summary reflects imported data"""
        # Get current summary
        response = requests.get(f"{BASE_URL}/api/participants/payment-summary", headers=commander_headers)
        assert response.status_code == 200
        data = response.json()
        
        # Verify last_import is populated if imports exist
        history = requests.get(f"{BASE_URL}/api/payment-imports", headers=commander_headers)
        if len(history.json()) > 0:
            assert data["last_import"] is not None, "last_import should be populated when imports exist"
            assert "filename" in data["last_import"], "last_import should have filename"
            assert "matched" in data["last_import"], "last_import should have matched count"
    
    def test_participants_have_payment_fields(self, commander_headers):
        """Participants in summary have payment-related fields"""
        response = requests.get(f"{BASE_URL}/api/participants/payment-summary", headers=commander_headers)
        assert response.status_code == 200
        data = response.json()
        
        if len(data["participants"]) > 0:
            # Check that payment fields exist
            p = data["participants"][0]
            assert "paid" in p, "Participant should have 'paid' field"
            assert "amount_paid" in p, "Participant should have 'amount_paid' field"
            assert "registration_status" in p, "Participant should have 'registration_status' field"
            assert "unit_approved" in p, "Participant should have 'unit_approved' field"
            assert "wing_approved" in p, "Participant should have 'wing_approved' field"


class TestBudgetEndpointsAccess:
    """Tests for budget endpoint access control"""
    
    def test_budget_requires_finance_access(self, cadre_headers):
        """Budget endpoint requires finance access"""
        response = requests.get(f"{BASE_URL}/api/budget", headers=cadre_headers)
        assert response.status_code == 403, f"Expected 403, got {response.status_code}"
    
    def test_budget_summary_requires_finance_access(self, cadre_headers):
        """Budget summary requires finance access"""
        response = requests.get(f"{BASE_URL}/api/budget/summary", headers=cadre_headers)
        assert response.status_code == 403, f"Expected 403, got {response.status_code}"
    
    def test_commander_can_access_budget(self, commander_headers):
        """Commander can access budget"""
        response = requests.get(f"{BASE_URL}/api/budget", headers=commander_headers)
        assert response.status_code == 200, f"Expected 200, got {response.status_code}"
    
    def test_commander_can_access_budget_summary(self, commander_headers):
        """Commander can access budget summary"""
        response = requests.get(f"{BASE_URL}/api/budget/summary", headers=commander_headers)
        assert response.status_code == 200, f"Expected 200, got {response.status_code}"


class TestRosterPaymentIndicators:
    """Tests for roster payment indicators (regression check)"""
    
    def test_participants_have_paid_field(self, commander_headers):
        """Participants endpoint returns paid field"""
        response = requests.get(f"{BASE_URL}/api/participants", headers=commander_headers)
        assert response.status_code == 200
        data = response.json()
        
        if len(data) > 0:
            # Check that paid field exists
            p = data[0]
            assert "paid" in p or "paid_in_full" in p, "Participant should have paid/paid_in_full field"
    
    def test_participants_payment_fields_after_import(self, commander_headers):
        """Participants retain payment fields after import"""
        # Get a participant from payment summary
        summary = requests.get(f"{BASE_URL}/api/participants/payment-summary", headers=commander_headers)
        assert summary.status_code == 200
        
        if len(summary.json()["participants"]) > 0:
            capid = summary.json()["participants"][0]["capid"]
            
            # Get full participant list and find this one
            participants = requests.get(f"{BASE_URL}/api/participants", headers=commander_headers)
            assert participants.status_code == 200
            
            matching = [p for p in participants.json() if p.get("capid") == capid]
            if len(matching) > 0:
                p = matching[0]
                # Verify payment fields exist
                assert "paid" in p or "paid_in_full" in p, "Participant should have paid field"


if __name__ == "__main__":
    pytest.main([__file__, "-v", "--tb=short"])
